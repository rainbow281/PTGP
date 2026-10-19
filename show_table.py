#!/usr/bin/env python3
"""Summarize experiment results."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


DATASET_START_ROW = {
    "imagenet": 3,
    "ffhq": 10,
    "cifar100": 17,
    "cifar10": 24,
}

ATTACK_LABEL = {
    "geiping": "IG",
    "yin": "GI",
    "ggl": "GGL",
    "gias": "GIAS",
    "gifd": "GIFD",
}

DEFENSE_OFFSET = {
    "none": 0,
    "noise": 1,
    "clipping": 2,
    "sparsi": 3,
    "soteria": 4,
    "censor": 5,
    "spectra": 6,
}

ATTACK_START_COLUMN = {
    "IG": 3,     # C
    "GI": 8,     # H
    "GGL": 13,   # M
    "GIAS": 18,  # R
    "GIFD": 23,  # W
}

METRICS = {
    "MSE": "test_mse",
    "LPIPS(VGG)": "LPIPS_VGG",
    "LPIPS(ALEX)": "LPIPS_ALEX",
    "PSNR": "psnr",
    "SSIM": "ssim",
}


def normalize_root(path: Path) -> Path:
    """Experiment support utility."""
    if not path.is_dir():
        raise FileNotFoundError(path)
    nested = path / "results"
    return nested if nested.is_dir() else path


def read_metrics(path: Path) -> pd.DataFrame | None:
    try:
        frame = pd.read_csv(path)
    except Exception:
        return None
    columns = list(METRICS.values())
    if not set(columns).issubset(frame.columns):
        return None
    values = frame[columns].apply(pd.to_numeric, errors="coerce")
    
    values = values.mask(values < 0)
    return values.dropna(how="any")


def collect_dataset(dataset: str, supplied_root: Path) -> dict:
    root = normalize_root(supplied_root)
    pattern = re.compile(
        rf"{re.escape(dataset)}_(geiping|yin|ggl|gias|gifd)_"
        rf"(none|noise|clipping|sparsi|soteria|censor|spectra)_seed",
        re.IGNORECASE,
    )
    selected: dict[tuple[str, str], tuple[int, float, pd.DataFrame, Path]] = {}

    for run_dir in root.iterdir():
        if not run_dir.is_dir():
            continue
        match = pattern.match(run_dir.name)
        if match is None:
            continue
        key = (match.group(1).lower(), match.group(2).lower())
        for csv_path in run_dir.rglob("*.csv"):
            values = read_metrics(csv_path)
            if values is None or values.empty:
                continue
            candidate = (len(values), csv_path.stat().st_mtime, values, csv_path)
            old = selected.get(key)
            if old is None or candidate[:2] > old[:2]:
                selected[key] = candidate

    summaries = {}
    for key, (count, _, values, source) in selected.items():
        item = {"n": count, "source": source}
        for display_name, column in METRICS.items():
            item[display_name] = {
                "mean": float(values[column].mean()),
                "std": float(values[column].std(ddof=1)) if count > 1 else None,
            }
        summaries[key] = item
    return summaries


def statistic_text(mean: float, std: float | None, precision: int) -> str:
    if std is None:
        return f"{mean:.{precision}f}"
    return f"{mean:.{precision}f}±{std:.{precision}f}"


def write_dataset(sheet, dataset: str, summaries: dict, precision: int) -> None:
    first_row = DATASET_START_ROW[dataset]
    for (attack_key, defense_key), summary in summaries.items():
        row = first_row + DEFENSE_OFFSET[defense_key]
        first_column = ATTACK_START_COLUMN[ATTACK_LABEL[attack_key]]
        for offset, metric in enumerate(METRICS):
            value = summary[metric]
            sheet.cell(row=row, column=first_column + offset).value = statistic_text(
                value["mean"], value["std"], precision
            )


def report_completeness(dataset: str, summaries: dict,
                        expected_samples: int) -> tuple[list, list]:
    expected = {
        (attack, defense)
        for attack in ATTACK_LABEL
        for defense in DEFENSE_OFFSET
    }
    missing = sorted(expected - set(summaries))
    incomplete = sorted(
        (attack, defense, item["n"])
        for (attack, defense), item in summaries.items()
        if item["n"] != expected_samples
    )
    print(f"[{dataset}] combinations: {len(summaries)}/35")
    for attack, defense, count in incomplete:
        print(f"  [INCOMPLETE] {attack}/{defense}: "
              f"n={count}, expected={expected_samples}")
    for attack, defense in missing:
        print(f"  [MISSING] {attack}/{defense}")
    return missing, incomplete


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--imagenet-root", type=Path)
    parser.add_argument("--ffhq-root", type=Path)
    parser.add_argument("--cifar100-root", type=Path)
    parser.add_argument("--cifar10-root", type=Path)
    parser.add_argument("--expected-samples", type=int, default=10)
    parser.add_argument("--precision", type=int, default=3)
    parser.add_argument("--require-complete", action="store_true")
    args = parser.parse_args()

    if not args.template.is_file():
        raise FileNotFoundError(args.template)

    roots = {
        "imagenet": args.imagenet_root,
        "ffhq": args.ffhq_root,
        "cifar100": args.cifar100_root,
        "cifar10": args.cifar10_root,
    }
    roots = {name: path for name, path in roots.items() if path is not None}
    if not roots:
        parser.error("At least one dataset root must be supplied")

    workbook = load_workbook(args.template)
    if "Sheet1" not in workbook.sheetnames:
        raise KeyError("Template does not contain worksheet 'Sheet1'")
    sheet = workbook["Sheet1"]

    has_problem = False
    for dataset, root in roots.items():
        summaries = collect_dataset(dataset, root)
        missing, incomplete = report_completeness(
            dataset, summaries, args.expected_samples
        )
        has_problem |= bool(missing or incomplete)
        write_dataset(sheet, dataset, summaries, args.precision)

    if args.require_complete and has_problem:
        raise RuntimeError("At least one supplied dataset is incomplete; no workbook written")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Saved: {args.output}")


if __name__ == "__main__":
    main()
