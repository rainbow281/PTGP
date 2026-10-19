#!/usr/bin/env python3
"""Aggregate four-dataset, five-attack, seven-defense reconstruction experiments."""
from __future__ import annotations

import argparse
import re
from itertools import product
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.lines import Line2D
import numpy as np
import pandas as pd


DATASETS = ("cifar10", "cifar100", "ffhq", "imagenet")
ATTACKS = ("geiping", "yin", "ggl", "gias", "gifd")
DEFENSES = ("none", "noise", "clipping", "sparsi", "soteria", "censor", "tine_dus")
METRICS = ("psnr", "ssim", "test_mse", "LPIPS_VGG", "LPIPS_ALEX", "feat_mse", "acc_after")

ATTACK_LABELS = {"geiping": "IG", "yin": "GI", "ggl": "GGL", "gias": "GIAS", "gifd": "GIFD"}
DEFENSE_LABELS = {
    "none": "None", "noise": "Noise", "clipping": "Clipping",
    "sparsi": "Sparsification", "soteria": "Soteria",
    "censor": "CENSOR", "tine_dus": "TINE-DUS",
}
SHORT_DEFENSE_LABELS = {
    "none": "None", "noise": "Noise", "clipping": "Clip",
    "sparsi": "Sparse", "soteria": "Soteria",
    "censor": "CENSOR", "tine_dus": "TINE-DUS",
}
METRIC_LABELS = {
    "psnr": "PSNR (dB)", "ssim": "SSIM", "test_mse": "MSE",
    "LPIPS_VGG": "LPIPS-VGG", "LPIPS_ALEX": "LPIPS-Alex",
    "feat_mse": "Feature MSE", "acc_after": "Accuracy",
}
plt.rcParams.update({
    "font.family": "serif",
    "font.serif": ["Times New Roman", "Liberation Serif", "Nimbus Roman", "STIXGeneral"],
    "mathtext.fontset": "stix", "pdf.fonttype": 42, "ps.fonttype": 42,
})


def parse_args() -> argparse.Namespace:
    base = Path("/mnt/sda/rainbow/yanch_gmm/comparison_results_9_24_1")
    parser = argparse.ArgumentParser()
    parser.add_argument("--cifar10-root", type=Path,
                        default=base / "cifar10_five_attack_full/results")
    parser.add_argument("--cifar100-root", type=Path,
                        default=base / "cifar100_five_attack_full/results")
    parser.add_argument("--ffhq-root", type=Path,
                        default=base / "ffhq_five_attack_full/results")
    parser.add_argument("--imagenet-root", type=Path,
                        default=base / "imagenet_five_attack_full/results")
    parser.add_argument("--output-dir", type=Path,
                        default=base / "four_dataset_summary")
    parser.add_argument("--expected-samples", type=int, default=10)
    parser.add_argument(
        "--allow-incomplete", action="store_true",
        help="Allow exploratory summaries with missing groups/samples. Formal tables must not use this flag.")
    return parser.parse_args()


def parse_run(path: Path, dataset: str):
    pattern = re.compile(
        rf"{dataset}_(geiping|yin|ggl|gias|gifd)_"
        r"(none|noise|clipping|compression|sparsi|soteria|censor|tine_dus)"
        r"_seed(\d+)", re.IGNORECASE)
    match = pattern.search(str(path))
    if not match:
        return None
    attack, defense, seed = match.groups()
    defense = "sparsi" if defense.lower() == "compression" else defense.lower()
    return attack.lower(), defense, int(seed)


def stage_priority(path: Path, attack: str) -> int:
    """Prefer the final cumulative GIFD stage if stage-specific metric CSVs exist."""
    if attack != "gifd":
        return 0
    stages = []
    for part in path.parts:
        match = re.fullmatch(r"best_first_(\d+)_layer", part, re.IGNORECASE)
        if match:
            stages.append(int(match.group(1)))
    return max(stages, default=0)


def experiment_dir(path: Path, dataset: str) -> Path | None:
    prefix = f"{dataset}_"
    return next((parent for parent in path.parents
                 if parent.name.lower().startswith(prefix)
                 and "_seed" in parent.name.lower()), None)


def read_accuracy(run_dir: Path | None, cache: dict) -> pd.DataFrame | None:
    if run_dir is None:
        return None
    key = str(run_dir)
    if key in cache:
        return cache[key]
    
    
    
    candidates = list(run_dir.rglob("table_*_epoch_accuracy.csv"))
    frames = []
    for path in candidates:
        try:
            frame = pd.read_csv(path)
        except Exception:
            continue
        if "target_id" in frame.columns and "acc_after" in frame.columns:
            if "accuracy_scope" in frame.columns:
                frame = frame.loc[
                    frame["accuracy_scope"] == "per_sample_independent_update"]
            frame = frame[["target_id", "acc_after"]].copy()
            frame["mtime"] = path.stat().st_mtime
            frames.append(frame)
    if not frames:
        cache[key] = None
    else:
        cache[key] = (pd.concat(frames).sort_values("mtime")
                      .drop_duplicates("target_id", keep="last")
                      .drop(columns="mtime"))
    return cache[key]


def load_dataset(dataset: str, root: Path) -> pd.DataFrame:
    if not root.is_dir():
        print(f"[MISSING ROOT] {dataset}: {root}")
        return pd.DataFrame()
    rows, accuracy_cache = [], {}
    required_reconstruction = set(METRICS[:-1])
    for csv_path in root.rglob("*.csv"):
        parsed = parse_run(csv_path, dataset)
        if parsed is None:
            continue
        attack, defense, seed = parsed
        try:
            frame = pd.read_csv(csv_path)
        except Exception as error:
            print(f"[SKIP unreadable] {csv_path}: {error}")
            continue
        if not required_reconstruction.issubset(frame.columns):
            continue
        selected = frame[list(METRICS[:-1])].apply(pd.to_numeric, errors="coerce")
        # Negative reconstruction metrics are failure sentinels, never real scores.
        for metric in METRICS[:-1]:
            selected.loc[selected[metric] < 0, metric] = np.nan
        selected["target_id"] = (pd.to_numeric(frame["target_id"], errors="coerce")
                                 if "target_id" in frame else np.arange(len(frame)))
        selected["dataset"] = dataset
        selected["attack"] = attack
        selected["defense"] = defense
        selected["seed"] = seed
        selected["source"] = str(csv_path)
        selected["source_mtime"] = csv_path.stat().st_mtime
        selected["stage"] = stage_priority(csv_path, attack)
        accuracy = read_accuracy(experiment_dir(csv_path, dataset), accuracy_cache)
        if accuracy is None:
            selected["acc_after"] = np.nan
        else:
            selected = selected.merge(accuracy, on="target_id", how="left")
        rows.append(selected)
    if not rows:
        print(f"[NO METRICS] {dataset}: {root}")
        return pd.DataFrame()
    result = pd.concat(rows, ignore_index=True)
    result = result.sort_values(["stage", "source_mtime"]).drop_duplicates(
        ["dataset", "attack", "defense", "seed", "target_id"], keep="last")
    return result


def completeness_table(results: pd.DataFrame, expected_samples: int) -> pd.DataFrame:
    grid = pd.DataFrame(product(DATASETS, ATTACKS, DEFENSES),
                        columns=["dataset", "attack", "defense"])
    if results.empty:
        grid["n_unique_samples"] = 0
    else:
        counts = (results.groupby(["dataset", "attack", "defense"])
                  .target_id.nunique().rename("n_unique_samples").reset_index())
        grid = grid.merge(counts, how="left").fillna({"n_unique_samples": 0})
    grid["n_unique_samples"] = grid.n_unique_samples.astype(int)
    if not results.empty:
        valid_counts = (results.groupby(["dataset", "attack", "defense"])[list(METRICS)]
                        .count().add_suffix("_valid").reset_index())
        grid = grid.merge(valid_counts, how="left").fillna(0)
    for metric in METRICS:
        column = f"{metric}_valid"
        if column not in grid:
            grid[column] = 0
        grid[column] = grid[column].astype(int)
    grid["expected_samples"] = expected_samples
    grid["complete"] = (grid.n_unique_samples.eq(expected_samples)
                        & grid[[f"{metric}_valid" for metric in METRICS]]
                        .eq(expected_samples).all(axis=1))
    return grid


def summarize(results: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    keys = ["dataset", "attack", "defense", "metric"]
    long = results.melt(
        id_vars=["dataset", "attack", "defense", "seed", "target_id"],
        value_vars=list(METRICS), var_name="metric", value_name="value")
    summary = (long.groupby(keys, dropna=False).value
               .agg(n="count", mean="mean", std="std", variance="var",
                    median="median", q1=lambda x: x.quantile(.25),
                    q3=lambda x: x.quantile(.75), minimum="min", maximum="max")
               .reset_index())
    summary["mean_std"] = summary.apply(
        lambda row: (f"{row['mean']:.6f} ± {row['std']:.6f}"
                     if pd.notna(row["std"]) else f"{row['mean']:.6f} ± NA"), axis=1)
    readable = summary.pivot(index=["dataset", "attack", "defense"],
                             columns="metric", values="mean_std").reset_index()
    # Numeric wide table: every metric has explicit mean/std/variance/n columns.
    numeric = summary.pivot(
        index=["dataset", "attack", "defense"], columns="metric",
        values=["n", "mean", "std", "variance"]).reset_index()
    numeric.columns = [
        "_".join(str(item) for item in column if str(item))
        if isinstance(column, tuple) else str(column)
        for column in numeric.columns
    ]
    return summary, readable, numeric


def boxplot_pages(results: pd.DataFrame, output: Path) -> None:
    pdf_path = output / "four_dataset_metric_boxplots.pdf"
    count_rows = []
    with PdfPages(pdf_path) as pdf:
        for dataset in DATASETS:
            subset = results[results.dataset == dataset]
            if subset.empty:
                continue
            fig, axes = plt.subplots(len(METRICS), len(ATTACKS),
                                     figsize=(14.2, 17.5), squeeze=False)
            for row, metric in enumerate(METRICS):
                for col, attack in enumerate(ATTACKS):
                    ax = axes[row, col]
                    groups = [subset.loc[(subset.attack == attack)
                                         & (subset.defense == defense), metric]
                              .dropna().to_numpy() for defense in DEFENSES]
                    for defense, values in zip(DEFENSES, groups):
                        count_rows.append({"dataset": dataset, "attack": attack,
                                           "defense": defense, "metric": metric,
                                           "n": len(values)})
                    positions = np.arange(1, len(DEFENSES) + 1)
                    valid = [(position, defense, values)
                             for position, defense, values in zip(positions, DEFENSES, groups)
                             if len(values)]
                    if valid:
                        boxes = ax.boxplot(
                            [item[2] for item in valid],
                            positions=[item[0] for item in valid], widths=.62,
                            patch_artist=True, notch=False, whis=1.5,
                            showmeans=True, showfliers=True,
                            boxprops={"facecolor": "white", "edgecolor": "black",
                                      "linewidth": 1.0},
                            whiskerprops={"color": "black", "linewidth": .9},
                            capprops={"color": "black", "linewidth": .9},
                            medianprops={"color": "#C00000", "linewidth": 1.25},
                            meanprops={"marker": "D", "markersize": 3.1,
                                       "markerfacecolor": "black",
                                       "markeredgecolor": "black"},
                            flierprops={"marker": "o", "markersize": 2.5,
                                       "markerfacecolor": "none",
                                       "markeredgecolor": "black", "alpha": .7})
                    ax.set_xticks(positions,
                                  [f"{SHORT_DEFENSE_LABELS[d]}\nn={len(v)}"
                                   for d, v in zip(DEFENSES, groups)],
                                  rotation=55, ha="right", fontsize=6.5)
                    if row == 0:
                        ax.set_title(ATTACK_LABELS[attack], fontsize=11)
                    if col == 0:
                        ax.set_ylabel(METRIC_LABELS[metric], fontsize=9.5)
                    ax.grid(axis="y", alpha=.22, linewidth=.55)
                    ax.tick_params(axis="y", labelsize=7)
            fig.suptitle(dataset.upper(), fontsize=14, y=.997)
            legend = [
                Line2D([0], [0], color="#C00000", linewidth=1.4, label="Median"),
                Line2D([0], [0], marker="D", color="black", linewidth=0,
                       markersize=4, label="Mean"),
            ]
            fig.legend(handles=legend, loc="lower center", ncol=7,
                       bbox_to_anchor=(.5, .002), frameon=False, fontsize=8.5)
            fig.subplots_adjust(left=.07, right=.992, bottom=.045, top=.98,
                                wspace=.25, hspace=.55)
            pdf.savefig(fig)
            fig.savefig(output / f"{dataset}_metric_boxplots.pdf")
            fig.savefig(output / f"{dataset}_metric_boxplots.png", dpi=300)
            plt.close(fig)
    pd.DataFrame(count_rows).to_csv(output / "boxplot_sample_counts.csv", index=False)
    print(f"Boxplots: {pdf_path}")


def write_excel(results, summary, readable, numeric, completeness, output):
    path = output / "four_dataset_attack_defense_summary.xlsx"
    try:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            readable.to_excel(writer, sheet_name="mean_std", index=False)
            numeric.to_excel(writer, sheet_name="numeric_mean_std", index=False)
            summary.to_excel(writer, sheet_name="summary_long", index=False)
            completeness.to_excel(writer, sheet_name="completeness", index=False)
            results.to_excel(writer, sheet_name="sample_metrics", index=False)
        # Paper table follows the supplied reference: datasets/methods in the
        # first two columns and five attacks as horizontal metric groups.
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        workbook = load_workbook(path)
        sheet = workbook.create_sheet("paper_table", 0)
        metric_order = ["test_mse", "LPIPS_VGG", "LPIPS_ALEX", "psnr", "ssim"]
        metric_headers = ["MSE ↑", "LPIPS (VGG) ↑", "LPIPS (Alex) ↑", "PSNR ↓", "SSIM ↓"]
        attack_fills = ["D9EAF7", "E2F0D9", "FFF2CC", "FCE4D6", "E4DFEC"]
        dataset_labels = {"cifar10": "CIFAR-10", "cifar100": "CIFAR-100",
                          "ffhq": "FFHQ", "imagenet": "ImageNet"}
        lookup = summary.set_index(["dataset", "attack", "defense", "metric"])
        thin = Side(style="thin", color="7F7F7F")
        medium = Side(style="medium", color="3F3F3F")
        all_border = Border(left=thin, right=thin, top=thin, bottom=thin)

        sheet["A1"] = "Dataset"; sheet["B1"] = "Method"
        sheet.merge_cells("A1:A2"); sheet.merge_cells("B1:B2")
        for attack_index, attack in enumerate(ATTACKS):
            first_col = 3 + attack_index * len(metric_order)
            last_col = first_col + len(metric_order) - 1
            sheet.merge_cells(start_row=1, start_column=first_col,
                              end_row=1, end_column=last_col)
            attack_cell = sheet.cell(1, first_col, ATTACK_LABELS[attack])
            attack_cell.fill = PatternFill("solid", fgColor=attack_fills[attack_index])
            for offset, label in enumerate(metric_headers):
                cell = sheet.cell(2, first_col + offset, label)
                cell.fill = PatternFill("solid", fgColor=attack_fills[attack_index])

        row_number = 3
        for dataset in DATASETS:
            dataset_start = row_number
            # Compute rankings independently for each attack and metric.
            ranks = {}
            for attack in ATTACKS:
                for metric in metric_order:
                    candidates = []
                    for defense in DEFENSES:
                        key = (dataset, attack, defense, metric)
                        if key in lookup.index and pd.notna(lookup.loc[key, "mean"]):
                            candidates.append((defense, float(lookup.loc[key, "mean"])))
                    candidates.sort(key=lambda item: item[1],
                                    reverse=metric not in {"psnr", "ssim"})
                    distinct = []
                    for _, value in candidates:
                        if not distinct or not np.isclose(value, distinct[-1]):
                            distinct.append(value)
                    ranks[(attack, metric)] = {
                        defense: (1 if distinct and np.isclose(value, distinct[0])
                                  else 2 if len(distinct) > 1 and np.isclose(value, distinct[1])
                                  else 0)
                        for defense, value in candidates}
            for defense in DEFENSES:
                sheet.cell(row_number, 1, dataset_labels[dataset])
                sheet.cell(row_number, 2, DEFENSE_LABELS[defense])
                for attack_index, attack in enumerate(ATTACKS):
                    for metric_index, metric in enumerate(metric_order):
                        column = 3 + attack_index * len(metric_order) + metric_index
                        key = (dataset, attack, defense, metric)
                        cell = sheet.cell(row_number, column)
                        if key in lookup.index:
                            record = lookup.loc[key]
                            mean, std = record["mean"], record["std"]
                            cell.value = (f"{mean:.4f} ± {std:.4f}" if pd.notna(std)
                                          else f"{mean:.4f} ± NA")
                        rank = ranks[(attack, metric)].get(defense, 0)
                        if rank == 1:
                            cell.font = Font(name="Times New Roman", size=10, bold=True,
                                             color="000000")
                        elif rank == 2:
                            cell.font = Font(name="Times New Roman", size=10,
                                             color="0070C0")
                row_number += 1
            sheet.merge_cells(start_row=dataset_start, start_column=1,
                              end_row=row_number - 1, end_column=1)
            for column in range(1, 28):
                sheet.cell(row_number - 1, column).border = Border(
                    left=thin, right=thin, top=thin, bottom=medium)

        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row,
                                   min_col=1, max_col=27):
            for cell in row:
                if cell.font.name != "Times New Roman":
                    cell.font = Font(name="Times New Roman", size=10,
                                     bold=cell.font.bold, color="000000")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if cell.border == Border():
                    cell.border = all_border
        for column in ("A", "B"):
            sheet.column_dimensions[column].width = 16
        for column in range(3, 28):
            sheet.column_dimensions[get_column_letter(column)].width = 18
        sheet.row_dimensions[1].height = 25
        sheet.row_dimensions[2].height = 28
        for row in range(3, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 21
        sheet.freeze_panes = "C3"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_area = f"A1:AA{sheet.max_row}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        workbook.save(path)
    except ImportError:
        print("[WARNING] openpyxl is unavailable; CSV tables were still generated.")


def main() -> None:
    args = parse_args()
    roots = {
        "cifar10": args.cifar10_root, "cifar100": args.cifar100_root,
        "ffhq": args.ffhq_root, "imagenet": args.imagenet_root,
    }
    args.output_dir.mkdir(parents=True, exist_ok=True)
    frames = [load_dataset(dataset, root) for dataset, root in roots.items()]
    frames = [frame for frame in frames if not frame.empty]
    if not frames:
        raise SystemExit("No valid metric CSV files were found in the four result roots.")
    results = pd.concat(frames, ignore_index=True)
    results.to_csv(args.output_dir / "all_sample_metrics.csv", index=False)
    completeness = completeness_table(results, args.expected_samples)
    completeness.to_csv(args.output_dir / "experiment_completeness.csv", index=False)
    missing = completeness[~completeness.complete]
    missing.to_csv(args.output_dir / "missing_or_incomplete_groups.csv", index=False)
    if not missing.empty and not args.allow_incomplete:
        raise SystemExit(
            "Formal summary aborted: incomplete attack-defense groups were found. "
            "See missing_or_incomplete_groups.csv. Re-run experiments after applying "
            "the pandas>=2 CSV append fix; use --allow-incomplete only for debugging.")
    summary, readable, numeric = summarize(results)
    summary.to_csv(args.output_dir / "summary_long.csv", index=False)
    readable.to_csv(args.output_dir / "summary_mean_std.csv", index=False)
    numeric.to_csv(args.output_dir / "summary_mean_std_numeric.csv", index=False)
    write_excel(results, summary, readable, numeric, completeness, args.output_dir)
    boxplot_pages(results, args.output_dir)
    print(f"Observed groups: {len(completeness) - len(missing)}/{len(completeness)} complete")
    print(f"Tables and figures: {args.output_dir}")


if __name__ == "__main__":
    main()
